import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import json

# ------------------ Utility Helpers for Table/Cell Styling ------------------
def set_cell_background(cell, fill_hex):
    """Set cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for table cells in twentieths of a point (dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_with_color(doc, text, level, rgb_color):
    """Add a heading with specific styling."""
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = 'Segoe UI'
        r.font.color.rgb = rgb_color
    return h

def add_paragraph_with_spacing(doc, text="", space_after=8, bold=False, italic=False, font_size=Pt(10.5), color=RGBColor(51, 51, 51)):
    """Add standard body paragraph with customized properties."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = font_size
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p

def add_code_block(doc, code_text):
    """Add a shaded code block with fixed-width font."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F2F4F4") # Light grey fill
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border setting (thick grey/blue border)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="1A5276"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    # Add code lines
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(44, 62, 80)
    
    # Spacer paragraph below code block
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(8)

def add_bullet_point(doc, title, desc, rgb_color):
    """Add a styled bullet list item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    
    r_title = p.add_run(title + ": ")
    r_title.font.name = 'Segoe UI'
    r_title.font.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = rgb_color
    
    r_desc = p.add_run(desc)
    r_desc.font.name = 'Segoe UI'
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = RGBColor(80, 80, 80)
    return p

# ------------------ Main Document Generator ------------------
def build_handbook():
    doc = docx.Document()
    
    # 1. Page Margins Setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    COLOR_PRIMARY = RGBColor(26, 82, 118)    # Deep Blue (Corporate / Technical)
    COLOR_SECONDARY = RGBColor(93, 109, 126)  # Slate Grey
    COLOR_HIGHLIGHT = RGBColor(220, 118, 51)  # Orange
    COLOR_TEXT = RGBColor(51, 51, 51)
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    
    # =========================================================================
    # COVER PAGE
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(100)
    title_run = title_p.add_run("DocuConverter Pro")
    title_run.font.size = Pt(38)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(220)
    sub_run = subtitle_p.add_run("The Complete Developer Handbook & Reference Manual\nComprehensive Technical Documentation")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = COLOR_SECONDARY
    
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run("System Release: Version 1.2.0\nAuthor: Lead Architect & Development Team\nDate: June 30, 2026")
    meta_run.font.size = Pt(10.5)
    meta_run.font.color.rgb = COLOR_SECONDARY
    
    doc.add_page_break()
    
    # =========================================================================
    # SECTION 1: PROJECT OVERVIEW
    # =========================================================================
    add_heading_with_color(doc, "1. Project Overview & Architecture", 1, COLOR_PRIMARY)
    
    add_paragraph_with_spacing(doc, "1.1 What is DocuConverter Pro?", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_paragraph_with_spacing(doc, 
        "DocuConverter Pro is an enterprise-grade file utility and document conversion suite. The platform allows users "
        "to upload, convert, split, merge, compress, and edit various document types on-the-fly. The toolset is built as "
        "a Django web application combined with robust backend services including Microsoft Word/PowerPoint COM automation interfaces, "
        "PyMuPDF rendering scripts, and PDF-to-Excel analytical parsers. Unlike typical SaaS utilities that break styling, "
        "DocuConverter Pro is structured around high-fidelity document layout matching."
    )
    
    add_paragraph_with_spacing(doc, "1.2 Project Objective", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_paragraph_with_spacing(doc,
        "The project aims to solve the core challenges of cross-format file operations: preserving font layouts, table borders, "
        "and margins during document conversions. By using native Microsoft COM automation on Windows hosts, it ensures that "
        "DOCX and PPTX to PDF outputs are rendered identically to desktop exports, whilst employing open-source engines "
        "(PyMuPDF, pdf2docx, mammoth) for sub-page processing, compression, and visual editing."
    )
    
    add_paragraph_with_spacing(doc, "1.3 Features", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_bullet_point(doc, "PDF to Word", "Rebuilds PDF files into native DOCX formatted documents.", COLOR_PRIMARY)
    add_bullet_point(doc, "Word to PDF", "Executes high-fidelity PDF output using Microsoft Word COM automation.", COLOR_PRIMARY)
    add_bullet_point(doc, "PDF Merge & Split", "Extracts page selections or concatenates multiple PDF document segments.", COLOR_PRIMARY)
    add_bullet_point(doc, "PDF Compress", "Reduces file size by scaling, compressing, and downsampling graphic image bytes.", COLOR_PRIMARY)
    add_bullet_point(doc, "Word Editor", "Visual WYSIWYG editor simulating MS Word, exporting canvas HTML back to DOCX format.", COLOR_PRIMARY)
    add_bullet_point(doc, "PDF Editor", "Executes text replacements and highlights directly on existing PDF text streams.", COLOR_PRIMARY)
    add_bullet_point(doc, "Excel/PPT Utilities", "Converts spreadsheets and slides to PDFs and formats data structures.", COLOR_PRIMARY)
    
    add_paragraph_with_spacing(doc, "1.4 Technologies Used", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_bullet_point(doc, "Python 3.11+", "Core programming language running script definitions.", COLOR_PRIMARY)
    add_bullet_point(doc, "Django 6.0.6", "Web MVC framework for routing, settings, and form processing.", COLOR_PRIMARY)
    add_bullet_point(doc, "HTML5 / CSS3 / JavaScript", "Frontend visual dashboard, styling, theme toggle, and Word modal workspace.", COLOR_PRIMARY)
    add_bullet_point(doc, "SQLite3", "Default database engine for sessions, auth logs, and history tracking.", COLOR_PRIMARY)
    
    add_paragraph_with_spacing(doc, "1.5 Architecture Diagram", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_paragraph_with_spacing(doc, "Below is the structural flow diagram representing the web interface, application service layer, and system utilities:")
    
    add_code_block(doc, 
        "+---------------------------------------------------------------------------------+\n"
        "|                              WEB BROWSER CLIENT                                 |\n"
        "|          - User Dashboard UI     - Theme Switcher       - Word Modal            |\n"
        "+----------------------------------------+----------------------------------------+\n"
        "                                         | HTTP POST Upload / AJAX Requests\n"
        "                                         v\n"
        "+----------------------------------------+----------------------------------------+\n"
        "|                              DJANGO WEB APP                                     |\n"
        "|          - urls.py (Routing)     - views/ (Controllers) - settings.py (Config)  |\n"
        "+----------------------------------------+----------------------------------------+\n"
        "                                         | Service Layer Calls\n"
        "                                         v\n"
        "+----------------------------------------+----------------------------------------+\n"
        "|                            SERVICES & ENGINES LAYER                             |\n"
        "|  +-----------------------+  +-----------------------+  +---------------------+  |\n"
        "|  |   word_to_pdf.py      |  |   pdf_service.py      |  |   compression.py    |  |\n"
        "|  |   (win32com / Word)   |  |   (fitz / PyMuPDF)    |  |   (fitz / Pillow)   |  |\n"
        "|  +-----------------------+  +-----------------------+  +---------------------+  |\n"
        "|  +-----------------------+  +-----------------------+  +---------------------+  |\n"
        "|  |   pdf_to_word.py      |  |   editor.py           |  |   excel_service.py  |  |\n"
        "|  |   (pdf2docx)          |  |   (fitz Replace)      |  |   (pdfplumber)      |  |\n"
        "|  +-----------------------+  +-----------------------+  +---------------------+  |\n"
        "+----------------------------------------+----------------------------------------+\n"
        "                                         | Reads / Writes\n"
        "                                         v\n"
        "+----------------------------------------+----------------------------------------+\n"
        "|                             PERSISTENCE & DISK                                  |\n"
        "|          - media/ (Converted files)    - db.sqlite3 (Sessions / Settings)       |\n"
        "+---------------------------------------------------------------------------------+"
    )
    
    add_paragraph_with_spacing(doc, "1.6 Folder Structure", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_code_block(doc,
        "c:/WordToPDF_Web/\n"
        "│\n"
        "├── .venv/                         # Activated python virtual environment\n"
        "├── packages_guide.pdf             # System reference guide (PDF)\n"
        "└── wordpdf/                       # Core Django project root\n"
        "    ├── manage.py                  # Administrative tasks runner\n"
        "    ├── db.sqlite3                 # Local database storage\n"
        "    ├── media/                     # Uploads & converted files directory\n"
        "    ├── templates/                 # General workspace templates directory\n"
        "    │   └── includes/              # Sub-modal templates (sidebar, word_modal)\n"
        "    ├── pdfapp/                    # Core conversion logic app\n"
        "    │   ├── admin.py / models.py   # Admin registration & conversion metadata\n"
        "    │   ├── urls/                  # Dynamic URLs grouping (word, pdf, excel)\n"
        "    │   ├── views/                 # View controllers (dashboard, word_views)\n"
        "    │   ├── services/              # Core conversion engines\n"
        "    │   │   ├── word_com.py        # COM Automation handler\n"
        "    │   │   ├── compression.py     # PDF Image compressor\n"
        "    │   │   └── editor.py          # PDF Text replace service\n"
        "    │   └── utils/                 # Utility validators & helper scripts\n"
        "    └── wordpdf/                   # System settings, core routing and WSGI\n"
    )
    
    doc.add_page_break()
    
    # =========================================================================
    # SECTION 2: COMPLETE PACKAGE GUIDE
    # =========================================================================
    add_heading_with_color(doc, "2. Complete Package & Dependency Directory", 1, COLOR_PRIMARY)
    
    # We will document the key packages inside the codebase
    package_list = [
        {
            "name": "Django (6.0.6)",
            "purpose": "Core Web MVC Framework",
            "install": "pip install django",
            "doc": "https://docs.djangoproject.com/en/6.0/",
            "adv": "Built-in security features, complete session ORM, structure settings configurations, and rapid admin dashboards.",
            "disadv": "Monolithic architecture, high startup memory usage compared to Flask/FastAPI.",
            "alt": "Flask, FastAPI, Pyramid",
            "code": "from django.shortcuts import render\nfrom django.http import JsonResponse\n\ndef my_view(request):\n    return JsonResponse({'status': 'running'})",
            "internal": "Receives WSGI requests, triggers middleware sequences, processes URL paths, resolves patterns to views, processes templates, and yields response streams.",
            "files": "manage.py, wordpdf/settings.py, wordpdf/urls.py, pdfapp/views/*.py",
            "questions": "1. What is the role of settings.py?\n2. Explain CSRF protection middleware.\n3. How do Django sessions work?\n4. What are Django migrations?\n5. Explain class-based vs function-based views.",
            "errors": "CSRF verification failed (Ensure token is in form). ImproperlyConfigured (Review settings.py database engine)."
        },
        {
            "name": "PyMuPDF / fitz (1.27.2)",
            "purpose": "High-Performance PDF Operations",
            "install": "pip install PyMuPDF",
            "doc": "https://pymupdf.readthedocs.io/en/latest/",
            "adv": "Extremely fast text parsing, multi-page layout rendering, text replacement capability, and low memory consumption.",
            "disadv": "Native C++ bindings can cause compilation issues on non-standard architectures; GPL licensing rules.",
            "alt": "PyPDF2, PDFminer.six, pdfplumber",
            "code": "import fitz\ndoc = fitz.open('input.pdf')\nprint('Total Pages:', len(doc))\nfor page in doc:\n    print(page.get_text())",
            "internal": "Wraps the MuPDF library in C, parsing the internal cross-reference tables (xref) of PDF streams to retrieve, add, or edit text coordinates.",
            "files": "pdfapp/services/pdf_service.py, pdfapp/services/editor.py, pdfapp/services/compression.py",
            "questions": "1. How does PyMuPDF extract images?\n2. How to replace text in PyMuPDF?\n3. Explain fitz.open().\n4. What is page.get_drawings()?\n5. How do you compress a PDF using PyMuPDF?",
            "errors": "RuntimeError: fitz not installed (Ensure package PyMuPDF is installed). ValueError: invalid xref (File is corrupted)."
        },
        {
            "name": "python-docx (1.2.0)",
            "purpose": "Microsoft Word Document Creation & Parsing",
            "install": "pip install python-docx",
            "doc": "https://python-docx.readthedocs.io/en/latest/",
            "adv": "Allows programmatically creating Word files, setting headers/footers, generating tables, and managing styles without Office.",
            "disadv": "Does not support older .doc formats. Cannot render documents to visual formats (PDF/PNG) on its own.",
            "alt": "python-docx2txt, docx",
            "code": "import docx\ndoc = docx.Document()\ndoc.add_heading('Hello Word', 0)\ndoc.add_paragraph('This is programmatically generated.')\ndoc.save('test.docx')",
            "internal": "Unzips the DOCX openXML ZIP file, parses document.xml DOM tree, and modifies node components (w:p, w:r, w:t) before re-zipping.",
            "files": "pdfapp/views_legacy.py, pdfapp/views/word_views.py, scripts/generate_guide.py",
            "questions": "1. What is a run in python-docx?\n2. How do you add a table cell?\n3. Can python-docx read headers?\n4. How to format font sizes?\n5. How does python-docx write XML namespaces?",
            "errors": "PackageNotFoundError (Invalid template path). AttributeError: 'Paragraph' object has no attribute 'font' (Font belongs to runs)."
        },
        {
            "name": "pywin32 / win32com (312)",
            "purpose": "Windows COM Automation Client",
            "install": "pip install pywin32",
            "doc": "https://github.com/mhammond/pywin32",
            "adv": "Gives Python access to native Windows APIs and desktop software (Microsoft Word, Excel, PowerPoint) for 100% conversion accuracy.",
            "disadv": "Works strictly on Windows platforms. Requires Microsoft Office desktop licenses running on the host system.",
            "alt": "python-docx (for writing), libreoffice CLI (for conversions)",
            "code": "import win32com.client\nword = win32com.client.DispatchEx('Word.Application')\nword.Visible = False\ndoc = word.Documents.Open('C:\\\\test.docx')\ndoc.SaveAs2('C:\\\\test.pdf', FileFormat=17)\ndoc.Close()\nword.Quit()",
            "internal": "Acts as an OLE/COM Client interface calling Word's application class via the IDispatch automation interface.",
            "files": "pdfapp/services/word_com.py, pdfapp/services/word_to_pdf.py",
            "questions": "1. What is DispatchEx?\n2. How to close active Winword instances?\n3. Explain CoInitialize.\n4. Why does OLE automation fail in services?\n5. What is FileFormat=17?",
            "errors": "com_error: Member not found (Invalid API method). PyIDispatch::Invoke failed (Word is locked or showing alert)."
        },
        {
            "name": "pdf2docx (0.5.13)",
            "purpose": "PDF to DOCX Converter",
            "install": "pip install pdf2docx",
            "doc": "https://dothinking.github.io/pdf2docx/",
            "adv": "Converts PDF tables, columns, and styles back to Word structures without relying on Office APIs.",
            "disadv": "Slow on large documents with complex graphics; struggles with handwritten or non-OCR scanned PDF pages.",
            "alt": "Adobe Acrobat Conversion Service, PyMuPDF, pdfminer",
            "code": "from pdf2docx import Converter\ncv = Converter('test.pdf')\ncv.convert('test.docx', start=0, end=None)\ncv.close()",
            "internal": "Uses PyMuPDF to extract layout components (drawings, text segments, blocks) and runs algorithms to reassemble them into Word elements.",
            "files": "pdfapp/services/pdf_to_word.py, pdfapp/views_legacy.py",
            "questions": "1. How does pdf2docx identify tables?\n2. Can pdf2docx convert partial pages?\n3. Does it support scanned PDF files?\n4. What does the parser do internally?\n5. How to speed up conversion times?",
            "errors": "ImportError: no pdf2docx (Install PyMuPDF first). KeyError (Corrupted layout block map)."
        }
    ]
    
    for pkg in package_list:
        add_heading_with_color(doc, f"■ Package: {pkg['name']}", 2, COLOR_PRIMARY)
        
        # Details Table
        tbl = doc.add_table(rows=6, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        w_left = Inches(1.5)
        w_right = Inches(5.0)
        
        headers_info = [
            ("Primary Purpose", pkg["purpose"]),
            ("Installation Command", pkg["install"]),
            ("Official Documentation", pkg["doc"]),
            ("Key Advantages", pkg["adv"]),
            ("Key Disadvantages", pkg["disadv"]),
            ("Internal Workings", pkg["internal"])
        ]
        
        for idx, (label, val) in enumerate(headers_info):
            row = tbl.rows[idx]
            row.cells[0].text = label
            row.cells[1].text = val
            
            # Formats
            set_cell_background(row.cells[0], "EAEDED")
            set_cell_margins(row.cells[0], top=80, bottom=80, left=100, right=100)
            set_cell_margins(row.cells[1], top=80, bottom=80, left=100, right=100)
            
            row.cells[0].width = w_left
            row.cells[1].width = w_right
            
            # Bold labels
            for p in row.cells[0].paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(9.5)
            for p in row.cells[1].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    
        add_paragraph_with_spacing(doc).paragraph_format.space_after = Pt(4)
        
        # Code sample
        add_paragraph_with_spacing(doc, "Example Implementation Code:", bold=True, font_size=Pt(10), color=COLOR_SECONDARY)
        add_code_block(doc, pkg["code"])
        
        # Troubleshoots
        add_paragraph_with_spacing(doc, "Common Errors & Troubleshooting Checklist:", bold=True, font_size=Pt(10), color=COLOR_SECONDARY)
        add_bullet_point(doc, "Known Issue", pkg["errors"].split(".")[0], COLOR_HIGHLIGHT)
        add_bullet_point(doc, "Resolution", pkg["errors"].split(".")[1] if len(pkg["errors"].split(".")) > 1 else "Check package configuration", COLOR_PRIMARY)
        
        add_paragraph_with_spacing(doc).paragraph_format.space_after = Pt(12)
        
    doc.add_page_break()

    # =========================================================================
    # SECTION 3: PACKAGE INSTALLATION GUIDE
    # =========================================================================
    add_heading_with_color(doc, "3. System Package Installation Guide", 1, COLOR_PRIMARY)
    
    add_paragraph_with_spacing(doc, "Follow this complete setup checklist to run the application:")
    
    add_paragraph_with_spacing(doc, "Step 1: Setting up Virtual Environment", bold=True, font_size=Pt(11), color=COLOR_PRIMARY)
    add_code_block(doc, 
        "# Open powershell and navigate to the project directory\n"
        "cd c:\\WordToPDF_Web\n\n"
        "# Create venv named '.venv'\n"
        "python -m venv .venv\n\n"
        "# Activate venv\n"
        "# PowerShell:\n"
        ".venv\\Scripts\\Activate.ps1\n"
        "# CMD Prompt:\n"
        ".venv\\Scripts\\activate.bat\n"
        "# Linux / Mac:\n"
        "source .venv/bin/activate"
    )
    
    add_paragraph_with_spacing(doc, "Step 2: Installing Python Modules", bold=True, font_size=Pt(11), color=COLOR_PRIMARY)
    add_code_block(doc,
        "# Upgrade package installer pip to latest\n"
        "python -m pip install --upgrade pip\n\n"
        "# Install core dependencies\n"
        "pip install django pymupdf python-docx docxcompose mammoth htmldocx pdf2docx pdfplumber openpyxl pillow pywin32 opencv-python-headless\n\n"
        "# Generate requirements definition\n"
        "pip freeze > requirements.txt\n\n"
        "# Install on other machines from definition\n"
        "pip install -r requirements.txt"
    )
    
    add_paragraph_with_spacing(doc, "Step 3: Post-Install Setup for OLE/COM Automation", bold=True, font_size=Pt(11), color=COLOR_PRIMARY)
    add_code_block(doc,
        "# Run pywin32 system config helper to initialize registry keys\n"
        "python .venv\\Scripts\\pywin32_postinstall.py -install"
    )
    
    doc.add_page_break()

    # =========================================================================
    # SECTION 4: FUNCTION EXPLANATION
    # =========================================================================
    add_heading_with_color(doc, "4. Comprehensive Function-Level Reference", 1, COLOR_PRIMARY)
    
    # Load metadata.json to document the actual functions
    meta_data = {}
    try:
        with open("scratch/metadata.json", "r", encoding="utf-8") as f:
            meta_data = json.load(f)
    except Exception as e:
        print(f"Could not load metadata: {e}")
        
    doc_count = 0
    for file_path, data in sorted(meta_data.items()):
        # Only document core files to avoid huge bloated handbook size
        core_files = [
            "pdfapp\\services\\word_com.py",
            "pdfapp\\services\\word_to_pdf.py",
            "pdfapp\\services\\pdf_to_word.py",
            "pdfapp\\services\\pdf_service.py",
            "pdfapp\\services\\compression.py",
            "pdfapp\\services\\editor.py",
            "pdfapp\\views\\word_views.py",
            "pdfapp\\views\\pdf_views.py"
        ]
        if file_path not in core_files:
            continue
            
        funcs = data.get("functions", [])
        if not funcs:
            continue
            
        add_heading_with_color(doc, f"Module: {file_path}", 2, COLOR_PRIMARY)
        
        for f in funcs:
            doc_count += 1
            if doc_count > 25: # Keep it limited to top 25 functions to keep conversion times fast
                break
                
            p_func = doc.add_paragraph()
            p_func.paragraph_format.space_before = Pt(8)
            p_func.paragraph_format.space_after = Pt(2)
            p_func.paragraph_format.keep_with_next = True
            
            run_func = p_func.add_run(f"► function: {f['name']}({', '.join(f['args'])})")
            run_func.font.bold = True
            run_func.font.size = Pt(11)
            run_func.font.color.rgb = COLOR_HIGHLIGHT
            
            # Purpose
            purpose = f["docstring"].strip() if f["docstring"] else "Executes targeted logical operations for this module handler."
            add_bullet_point(doc, "Purpose", purpose, COLOR_PRIMARY)
            
            # Signature details
            args_str = ", ".join(f["args"]) if f["args"] else "None"
            add_bullet_point(doc, "Parameters", args_str, COLOR_PRIMARY)
            add_bullet_point(doc, "Return Value", str(f.get("return_type", "None")), COLOR_PRIMARY)
            
            # Execution flow description
            flow_text = f"1. Receives inputs ({args_str}). 2. Resolves parameters. 3. Calls package dependencies. 4. Returns output."
            add_bullet_point(doc, "Execution Flow", flow_text, COLOR_PRIMARY)
            
            add_paragraph_with_spacing(doc).paragraph_format.space_after = Pt(4)
            
        if doc_count > 25:
            break
            
    doc.add_page_break()

    # =========================================================================
    # SECTION 5: DJANGO CONFIGURATION & FLOW
    # =========================================================================
    add_heading_with_color(doc, "5. Django Project Structure & Configuration", 1, COLOR_PRIMARY)
    
    django_details = [
        ("settings.py", "Contains system configurations including installed apps list (pdfapp), database connections, middleware definitions (Session, CSRF), and media upload/download path setups."),
        ("urls.py", "Performs path routing. Resolves input request URLs to match corresponding view handlers in the app views subfolder."),
        ("views.py", "Processes input requests. Views read inputs, call conversion services, save files to media/, and render the HTML interface templates."),
        ("models.py", "Defines the database schema using Django ORM. The ConversionHistory model stores converted filenames and timestamps."),
        ("admin.py", "Registers the DB models in the administrative panel enabling dashboard CRUD operations."),
        ("templates/", "Stores the visual presentation files (home.html, index.html) with embedded Django templates variables and tags."),
        ("app_scripts.html & localStorage", "Contains the client-side JavaScript dashboard logic. Standard HTML5 localStorage is used for client preference preservation: 'activeToolId' tracks the last active tool segment across page navigation, 'themeMode' retains user Light/Dark mode choice, and 'recentDownloads' stores a serialized JSON log of successfully converted files to populate the user's local activity table. All data is processed entirely in-browser.")
    ]
    
    for filename, desc in django_details:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r_name = p.add_run(f"■ {filename}")
        r_name.font.bold = True
        r_name.font.size = Pt(11)
        r_name.font.color.rgb = COLOR_PRIMARY
        
        doc.add_paragraph(desc).paragraph_format.left_indent = Inches(0.2)
        
    doc.add_page_break()

    # =========================================================================
    # SECTION 6: SERVICE WORKFLOWS
    # =========================================================================
    add_heading_with_color(doc, "6. Package Workflows & Flowcharts", 1, COLOR_PRIMARY)
    
    add_paragraph_with_spacing(doc, "6.1 Word to PDF Conversion Workflow", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_code_block(doc,
        "User uploads Word Document (.docx)\n"
        "             │\n"
        "             v\n"
        "Django handles upload -> saves temporarily in media/\n"
        "             │\n"
        "             v\n"
        "Invokes safe_docx_to_pdf(docx_path) service\n"
        "             │\n"
        "             v\n"
        "Loads win32com.client wrapper -> CoInitialize()\n"
        "             │\n"
        "             v\n"
        "Dispatches background 'Word.Application' process\n"
        "             │\n"
        "             v\n"
        "Opens document -> ExportAsFixedFormat(ExportFormat=PDF)\n"
        "             │\n"
        "             v\n"
        "Cleans up COM handles -> returns success response\n"
        "             │\n"
        "             v\n"
        "User receives file download link"
    )
    
    add_paragraph_with_spacing(doc, "6.2 PDF to Word Conversion Workflow", bold=True, font_size=Pt(12), color=COLOR_PRIMARY)
    add_code_block(doc,
        "User uploads PDF File (.pdf)\n"
        "             │\n"
        "             v\n"
        "Django receives PDF -> invokes pdf_to_word_service\n"
        "             │\n"
        "             v\n"
        "Service tries Word COM engine / pdf2docx fallback\n"
        "             │\n"
        "             v\n"
        "pdf2docx parses pages and constructs layout grids\n"
        "             │\n"
        "             v\n"
        "Reassembles paragraphs, fonts, images, and tables\n"
        "             │\n"
        "             v\n"
        "Saves as native DOCX structure -> persists to disk\n"
        "             │\n"
        "             v\n"
        "Frontend triggers celebratory animation and download"
    )
    
    doc.add_page_break()

    # =========================================================================
    # SECTION 7: COMMON COMMANDS REFERENCE
    # =========================================================================
    add_heading_with_color(doc, "7. Command-Line Reference Guide", 1, COLOR_PRIMARY)
    
    commands = [
        ("python manage.py runserver", 
         "Starts the Django local development server.", 
         "python manage.py runserver 8000",
         "Watching for file changes...\nPerforming system checks...\nStarting development server at http://127.0.0.1:8000/"),
         
        ("python manage.py migrate", 
         "Applies pending database migrations to db.sqlite3.", 
         "python manage.py migrate",
         "Operations to perform:\n  Apply all migrations: admin, auth, contenttypes, sessions\nRunning migrations:\n  Applying auth.0001_initial... OK"),
         
        ("python manage.py makemigrations", 
         "Scans app models.py files and generates migrations code.", 
         "python manage.py makemigrations",
         "Migrations for 'pdfapp':\n  pdfapp/migrations/0001_initial.py\n    - Create model ConversionHistory"),
         
        ("pip install <package>", 
         "Downloads and installs python packages from PyPI.", 
         "pip install pymupdf python-docx",
         "Downloading PyMuPDF-1.27.2.tar.gz...\nInstalling collected packages: PyMuPDF\nSuccessfully installed PyMuPDF-1.27.2"),
         
        ("pip freeze", 
         "Outputs all installed package versions to standard format.", 
         "pip freeze > requirements.txt",
         "django==6.0.6\npymupdf==1.27.2\npython-docx==1.2.0")
    ]
    
    for title, desc, syntax, output in commands:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run(f"■ Command: {title}")
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_PRIMARY
        
        doc.add_paragraph(desc).paragraph_format.space_after = Pt(4)
        
        add_paragraph_with_spacing(doc, f"Syntax: {syntax}", bold=True, font_size=Pt(9.5), color=COLOR_SECONDARY)
        add_code_block(doc, output)
        
    doc.add_page_break()

    # =========================================================================
    # SECTION 8: INTERVIEW QUESTIONS & PREPARATION
    # =========================================================================
    add_heading_with_color(doc, "8. Developer Technical Interview Guide", 1, COLOR_PRIMARY)
    
    questions = [
        ("Q1: What is COM automation and how does Django use it?",
         "Answer: COM (Component Object Model) is a binary-interface standard for software component automation. "
         "In this project, Django uses the pywin32 library to dispatch OLE-automation calls to Microsoft Word and PowerPoint, "
         "manipulating them in headless mode to convert files to PDFs with absolute layout fidelity."),
         
        ("Q2: Why is pythoncom.CoInitialize() required in views?",
         "Answer: When calling COM interface objects from within a multithreaded environment (like a Django WSGI server), "
         "each thread must initialize its COM apartment runtime space. Failing to run CoInitialize() leads to OLE thread errors."),
         
        ("Q3: How does PyMuPDF extract text, and how does it differ from pdfminer?",
         "Answer: PyMuPDF is a fast C-wrapped engine leveraging MuPDF, providing page text parsing in milliseconds. "
         "Pdfminer is written in pure Python and is much slower, although it offers more detailed font metric metadata structures."),
         
        ("Q4: What happens if Microsoft Word is not installed on the system?",
         "Answer: The application catches the OLE Dispatch Exception and raises a ConversionError. The system then redirects the "
         "operation or triggers fallback packages (like pdf2docx or mammoth) depending on the context."),
         
        ("Q5: Explain the structure of the DOCX format.",
         "Answer: A DOCX file is a zipped package containing OpenXML files. The main text content resides in word/document.xml, "
         "while styles, document properties, settings, footers, and headers reside in separate XML sheets linked by relationships."),
         
        ("Q6: How does pdf2docx reassemble tables from PDF streams?",
         "Answer: It extracts horizontal and vertical line drawings, analyzes their intersections to form cell coordinate bounds, "
         "groups overlapping text characters inside those bounds, and constructs docx table cell components in openXML.")
    ]
    
    for q, a in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        
        r_q = p_q.add_run(q)
        r_q.font.bold = True
        r_q.font.size = Pt(11)
        r_q.font.color.rgb = COLOR_PRIMARY
        
        p_a = doc.add_paragraph(a)
        p_a.paragraph_format.space_after = Pt(10)
        p_a.paragraph_format.left_indent = Inches(0.2)
        p_a.paragraph_format.line_spacing = 1.15
        
    doc.add_page_break()

    # =========================================================================
    # SECTION 9: ERRORS & TROUBLESHOOTING
    # =========================================================================
    add_heading_with_color(doc, "9. System Errors & Troubleshooting Directory", 1, COLOR_PRIMARY)
    
    errors = [
        ("1. com_error: (-2147221008, 'CoInitialize has not been called', None, None)",
         "Reason: The calling thread attempted to initialize a COM object (like Word) without initializing the thread COM apartment.",
         "Solution: Add 'import pythoncom; pythoncom.CoInitialize()' before dispatching win32com client calls, and run 'pythoncom.CoUninitialize()' in a finally block."),
         
        ("2. PermissionError: [Errno 13] Permission denied on media/ output file",
         "Reason: The target output file is locked by another running background thread, web request, or active MS Word process.",
         "Solution: Kill any stale WINWORD.EXE processes using Task Manager or run 'taskkill /IM WINWORD.EXE /F' in CMD."),
         
        ("3. CSRF token missing or incorrect in POST forms",
         "Reason: The user is submitting file conversion uploads without sending the cryptographic cross-site request forgery token.",
         "Solution: Ensure that '{% csrf_token %}' is embedded inside the form tag in the template HTML.")
    ]
    
    for title, reason, sol in errors:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run(f"■ Error: {title}")
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_HIGHLIGHT
        
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_after = Pt(2)
        p_r.paragraph_format.left_indent = Inches(0.2)
        p_r.add_run("Reason: ").font.bold = True
        p_r.add_run(reason)
        
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_after = Pt(10)
        p_s.paragraph_format.left_indent = Inches(0.2)
        p_s.add_run("Solution: ").font.bold = True
        p_s.add_run(sol)
        
    doc.add_page_break()

    # =========================================================================
    # SECTION 10: CODING BEST PRACTICES
    # =========================================================================
    add_heading_with_color(doc, "10. Development Coding Standards & Practices", 1, COLOR_PRIMARY)
    
    add_bullet_point(doc, "Code Quality", "Write scripts conforming to PEP 8 standards. Document complex services and exceptions.", COLOR_PRIMARY)
    add_bullet_point(doc, "Security Protocol", "Always validate input formats and verify file sizes before sending uploads to parser engines.", COLOR_PRIMARY)
    add_bullet_point(doc, "Thread Resource Control", "Always close document handles and call application.Quit() in COM blocks to avoid locking CPU resources.", COLOR_PRIMARY)
    add_bullet_point(doc, "Database Cleanups", "Periodically purge expired session data using the 'python manage.py clearsessions' utility.", COLOR_PRIMARY)
    
    doc.add_page_break()

    # =========================================================================
    # SECTION 11: SUMMARY COMPARISONS
    # =========================================================================
    add_heading_with_color(doc, "11. Package Comparison Summary", 1, COLOR_PRIMARY)
    
    sum_tbl = doc.add_table(rows=1, cols=4)
    sum_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_tbl.autofit = False
    
    hdr_sum = sum_tbl.rows[0].cells
    widths_sum = [Inches(1.8), Inches(1.8), Inches(1.4), Inches(1.5)]
    headers_sum = ["Package Name", "Primary Use", "Performance", "Best Used For"]
    
    for i, title in enumerate(headers_sum):
        hdr_sum[i].text = title
        hdr_sum[i].width = widths_sum[i]
        set_cell_background(hdr_sum[i], "1A5276")
        set_cell_margins(hdr_sum[i], top=100, bottom=100, left=120, right=120)
        for p in hdr_sum[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)
                
    comparisons = [
        ("Django", "Web Framework Router", "Medium", "Web Routing & DB"),
        ("PyMuPDF / fitz", "PDF Engine", "Fast", "Extract / Split / Merge"),
        ("python-docx", "DOCX Writer", "Fast", "Word Layout Builder"),
        ("pywin32 / COM", "High-Fidelity PDF Export", "Slow", "DOCX/PPTX to PDF"),
        ("pdf2docx", "PDF to DOCX Rebuilder", "Medium", "Word Exporting"),
        ("openpyxl", "Excel Processor", "Medium", "Spreadsheets Writing")
    ]
    
    for idx, (name, use, perf, best) in enumerate(comparisons):
        row = sum_tbl.add_row()
        cells = row.cells
        
        cells[0].text = name
        cells[1].text = use
        cells[2].text = perf
        cells[3].text = best
        
        bg_hex = "F4F6F7" if idx % 2 == 1 else "FFFFFF"
        for col_idx in range(4):
            cells[col_idx].width = widths_sum[col_idx]
            set_cell_background(cells[col_idx], bg_hex)
            set_cell_margins(cells[col_idx], top=80, bottom=80, left=120, right=120)
            for p in cells[col_idx].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.0)
                    if col_idx == 0:
                        r.font.bold = True
                        r.font.color.rgb = COLOR_PRIMARY
                        
    # Save the final file
    docx_path = os.path.abspath("media/DocuConverter_Pro_Handbook.docx")
    doc.save(docx_path)
    print(f"Developer handbook saved as DOCX: {docx_path}")
    return docx_path

def convert_to_pdf(docx_path):
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    
    import pythoncom
    import win32com.client
    
    pythoncom.CoInitialize()
    word = None
    try:
        print("Launching Microsoft Word COM to export developer handbook PDF...")
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        
        doc = word.Documents.Open(FileName=docx_path, ReadOnly=True, AddToRecentFiles=False)
        doc.ExportAsFixedFormat(
            OutputFileName=pdf_path,
            ExportFormat=17, # WD_EXPORT_FORMAT_PDF
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=False
        )
        doc.Close(SaveChanges=False)
        print(f"Developer handbook PDF exported: {pdf_path}")
    except Exception as e:
        print(f"COM conversion failed: {e}")
        try:
            from docx2pdf import convert
            convert(docx_path, pdf_path)
            print(f"PDF exported via docx2pdf fallback: {pdf_path}")
        except Exception as f_err:
            print(f"Fallback also failed: {f_err}")
            raise e
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()
    return pdf_path

if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.abspath(os.path.join(script_dir, ".."))
    os.chdir(project_root)
    
    d_path = build_handbook()
    p_path = convert_to_pdf(d_path)
    print(f"Conversion complete! PDF path: {p_path}")
